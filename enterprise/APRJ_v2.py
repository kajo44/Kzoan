import hashlib
import json
import logging
import os
import random
import shutil
import sqlite3
import sys
import threading
import traceback
import uuid
import zipfile
from datetime import datetime
from pathlib import Path, PurePosixPath

import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from PIL import Image, ImageTk
import openpyxl
from openpyxl.styles import Alignment as XlAlignment, Border, Font as XlFont, PatternFill, Side

import ai_model
import database
import pkrms_utils

APP_NAME = "APRJ"
APP_ROOT = (
    Path(sys.executable).resolve().parent
    if getattr(sys, "frozen", False)
    else Path(__file__).resolve().parent
)
APP_LOGS = APP_ROOT / "logs"
APP_ERROR_LOG = APP_LOGS / "application_errors.log"
SURVEY_AI_ROOT = APP_ROOT / "survey_ai"
SURVEY_AI_DATASET = SURVEY_AI_ROOT / "dataset"
SURVEY_AI_IMAGES = SURVEY_AI_DATASET / "images"
SURVEY_AI_MANIFEST = SURVEY_AI_DATASET / "labels.json"
SURVEY_AI_RUNS = SURVEY_AI_ROOT / "training_runs"
SURVEY_AI_MODELS = SURVEY_AI_ROOT / "models"
SURVEY_AI_MODEL = SURVEY_AI_MODELS / "survey_multi_task_model.pth"
SURVEY_AI_CACHE = SURVEY_AI_ROOT / "cache"
SURVEY_DB_PATH = APP_ROOT / "data" / "survey.sqlite3"


def log_application_error(context, error_type, error, error_traceback):
    try:
        APP_LOGS.mkdir(parents=True, exist_ok=True)
        logger = logging.getLogger(APP_NAME)
        if not logger.handlers:
            handler = logging.FileHandler(APP_ERROR_LOG, encoding="utf-8")
            handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
            logger.addHandler(handler)
            logger.setLevel(logging.ERROR)
        logger.error(
            "%s: %s",
            context,
            error,
            exc_info=(error_type, error, error_traceback),
        )
    except OSError as log_error:
        sys.stderr.write(f"{context}: {error}\nGagal menulis log {APP_ERROR_LOG}: {log_error}\n")


def survey_task_options():
    return {
        "bahu_kiri_tipe": ["Tidak Ada", "Beton", "Aspal", "Lapen", "Batu Kali", "Kerikil", "Tanah"],
        "bahu_kiri_kondisi": ["Tidak Ada", "Baik/Rata", "Bekas roda", "Erosi"],
        "bahu_kanan_tipe": ["Tidak Ada", "Beton", "Aspal", "Lapen", "Batu Kali", "Kerikil", "Tanah"],
        "bahu_kanan_kondisi": ["Tidak Ada", "Baik/Rata", "Bekas roda", "Erosi"],
        "drainase_kiri_tipe": ["Tidak Ada", "Tak perlu", "Tanah", "Pasang Batu Terbuka", "Pasang Batu Tertutup"],
        "drainase_kiri_kondisi": ["Tidak Ada", "Tersumbat", "Bersih", "Diperlukan Pasangan", "Diperlukan Drainase", "Erosi"],
        "drainase_kanan_tipe": ["Tidak Ada", "Tak perlu", "Tanah", "Pasang Batu Terbuka", "Pasang Batu Tertutup"],
        "drainase_kanan_kondisi": ["Tidak Ada", "Tersumbat", "Bersih", "Diperlukan Pasangan", "Diperlukan Drainase", "Erosi"],
        "perkerasan_tipe": ["Aspal", "Beton", "Blok Beton", "Lapen", "Batu Kali", "Kerikil", "Tanah"],
        "perkerasan_kondisi": ["BAIK", "SEDANG", "RUSAK RINGAN", "RUSAK BERAT"],
    }


SURVEY_TASKS = survey_task_options()
SURVEY_WIDTH_FIELDS = (
    "lebar_bahu_kiri",
    "lebar_drainase_kiri",
    "lebar_perkerasan",
    "lebar_bahu_kanan",
    "lebar_drainase_kanan",
)
SURVEY_AI_CLASS_OUTPUTS = tuple(f"ai_{task}" for task in SURVEY_TASKS)
SURVEY_AI_CONFIDENCE_OUTPUTS = tuple(f"ai_{task}_confidence" for task in SURVEY_TASKS)
SURVEY_AI_WIDTH_OUTPUTS = tuple(f"ai_{field}_m" for field in SURVEY_WIDTH_FIELDS)


def create_survey_model(torch_module, task_classes):
    class SurveyMultiTaskNet(torch_module.nn.Module):
        def __init__(self):
            super().__init__()
            self.features = torch_module.nn.Sequential(
                torch_module.nn.Conv2d(3, 32, 3, padding=1),
                torch_module.nn.ReLU(),
                torch_module.nn.MaxPool2d(2),
                torch_module.nn.Conv2d(32, 64, 3, padding=1),
                torch_module.nn.ReLU(),
                torch_module.nn.MaxPool2d(2),
                torch_module.nn.Conv2d(64, 128, 3, padding=1),
                torch_module.nn.ReLU(),
                torch_module.nn.AdaptiveAvgPool2d(1),
            )
            self.heads = torch_module.nn.ModuleDict(
                {key: torch_module.nn.Linear(128, len(values)) for key, values in task_classes.items()}
            )
            self.width_head = torch_module.nn.Linear(128, len(SURVEY_WIDTH_FIELDS))

        def forward(self, images):
            features = self.features(images).flatten(1)
            return {
                "classes": {key: head(features) for key, head in self.heads.items()},
                "widths": self.width_head(features),
            }

    return SurveyMultiTaskNet()



class SQLiteSurveyStore:
    """Penyimpanan survei lokal persisten berbasis SQLite."""

    def __init__(self, db_path):
        self.db_path = Path(db_path)
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        with sqlite3.connect(str(self.db_path), timeout=30) as conn:
            conn.execute("""
                CREATE TABLE IF NOT EXISTS survey_records (
                    photo_key TEXT PRIMARY KEY,
                    folder_source TEXT NOT NULL DEFAULT '',
                    survey_data TEXT NOT NULL,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
            """)

    def load_all(self):
        with sqlite3.connect(str(self.db_path), timeout=30) as conn:
            rows = conn.execute("SELECT photo_key, survey_data FROM survey_records").fetchall()
        result = {}
        for key, payload in rows:
            try:
                data = json.loads(payload)
                if isinstance(data, dict):
                    result[key] = data
            except (TypeError, json.JSONDecodeError):
                continue
        return result

    def save(self, photo_key, data):
        payload = json.dumps(data, ensure_ascii=False)
        with sqlite3.connect(str(self.db_path), timeout=30) as conn:
            conn.execute("""
                INSERT INTO survey_records(photo_key, folder_source, survey_data, updated_at)
                VALUES (?, ?, ?, CURRENT_TIMESTAMP)
                ON CONFLICT(photo_key) DO UPDATE SET
                    folder_source=excluded.folder_source,
                    survey_data=excluded.survey_data,
                    updated_at=CURRENT_TIMESTAMP
            """, (str(photo_key), str(data.get("folder", "")), payload))


class RoadSurveyApp:
    def __init__(self, root):
        self.root = root
        self.root.title(f"{APP_NAME} - Aplikasi Survei Jalan")
        self.root.report_callback_exception = self.handle_tk_callback_exception
        
        icon_path = APP_ROOT / "APRJ.ico"
        if icon_path.is_file():
            try:
                self.root.iconbitmap(str(icon_path))
            except tk.TclError as error:
                log_application_error("Memuat ikon aplikasi", type(error), error, error.__traceback__)
        
        self.root.geometry("1290x700")
        self.root.configure(bg="#eaf2f8")
        
        for directory in (SURVEY_AI_IMAGES, SURVEY_AI_RUNS, SURVEY_AI_MODELS, SURVEY_AI_CACHE):
            directory.mkdir(parents=True, exist_ok=True)
        
        self.photo_folder = ""
        self.photo_list = []
        self.current_photo_idx = -1
        self.survey_store = SQLiteSurveyStore(SURVEY_DB_PATH)
        self.survey_data = {}
        self._stored_survey_records = self.survey_store.load_all()
        
        self.ai_mode_enabled = False
        self.ai_resources = None
        self.ai_loaded_device = None
        self.ai_predictions = {}
        self.ai_status = tk.StringVar(value="AI nonaktif")
        self.survey_ai_model = None
        self.survey_ai_checkpoint = None
        self.ai_device_var = tk.StringVar(value="Otomatis")
        self.ai_batch_var = tk.StringVar(value="4")
        self.ai_epochs_var = tk.StringVar(value="10")
        self.ai_training_mode_var = tk.StringVar(
            value="Lanjutkan model aktif" if SURVEY_AI_MODEL.is_file() else "Latih ulang dari awal"
        )
        self.ai_training_status = tk.StringVar(value="Pilih export survei Excel untuk memulai.")
        self.staged_training_records = []
        self.training_running = False
        self.training_cancel_event = threading.Event()
        self.training_records = self.load_training_manifest()
        
        self.original_image = None
        self.zoom_scale = 1.0
        self.img_canvas_obj = None
        
        self.options_config = survey_task_options()
        self.buttons_ref = {}
        
        self.var_lebar_bahu_kiri = tk.StringVar(value="0")
        self.var_lebar_drainase_kiri = tk.StringVar(value="0")
        self.var_lebar_bahu_kanan = tk.StringVar(value="0")
        self.var_lebar_drainase_kanan = tk.StringVar(value="0")
        self.var_lebar_perkerasan = tk.StringVar(value="0")
        
        self.current_selections = {key: "" for key in self.options_config.keys()}
        
        self.build_ui()
        self.root.protocol("WM_DELETE_WINDOW", self.close_application)

    def photo_storage_key(self, filename):
        """Kunci stabil agar nama foto sama dari folder berbeda tidak bertabrakan."""
        return os.path.normcase(os.path.abspath(os.path.join(self.photo_folder, filename)))

    def close_application(self):
        try:
            self.root.destroy()
        except tk.TclError:
            pass

    def handle_tk_callback_exception(self, error_type, error, error_traceback):
        log_application_error("Kesalahan antarmuka", error_type, error, error_traceback)
        messagebox.showerror(
            "Terjadi kesalahan",
            f"{error_type.__name__}: {error}\n\nDetail tersimpan di:\n{APP_ERROR_LOG}",
            parent=self.root,
        )

    def show_help_and_diagnostics(self):
        try:
            import road_condition_gui as ai
            torch = ai.torch
            torch_status = f"PyTorch {torch.__version__}"
            cuda_status = (
                f"CUDA tersedia ({torch.cuda.get_device_name(0)})"
                if torch.cuda.is_available()
                else "CUDA tidak tersedia; AI memakai CPU"
            )
        except Exception:
            torch_status = "Runtime PyTorch eksternal tidak aktif/tidak dimuat"
            cuda_status = "Mode survei dasar dan multi-task mandiri"

        messagebox.showinfo(
            f"Panduan dan Diagnostik {APP_NAME}",
            f"Lokasi aplikasi:\n{APP_ROOT}\n\n"
            f"{torch_status}\n{cuda_status}\n\n"
            "Alur penggunaan: Pilih folder foto -> Muat Foto -> Isi/Koreksi Penilaian -> Ekspor Excel.\n"
            "Shortcut Keyboard: Panah Kiri (Sebelumnya), Panah Kanan (Simpan & Selanjutnya), Ctrl+R (Reset).\n\n"
            f"Log Kesalahan: {APP_ERROR_LOG}",
            parent=self.root,
        )

    def build_ui(self):
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill=tk.BOTH, expand=True)
        
        self.survey_tab = tk.Frame(self.notebook, bg="#eaf2f8")
        self.ai_tab = ttk.Frame(self.notebook, padding=12)
        self.pkrms_tab = ttk.Frame(self.notebook, padding=12)
        self.online_tab = ttk.Frame(self.notebook, padding=12)
        
        self.notebook.add(self.survey_tab, text="Survei Jalan")
        self.notebook.add(self.ai_tab, text="Rekayasa AI")
        self.notebook.add(self.pkrms_tab, text="PKRMS")
        self.notebook.add(self.online_tab, text="Integrasi Online")
        self.build_pkrms_tab()
        self.build_online_tab()

        # Top Control Frame
        top_frame = tk.Frame(self.survey_tab, bg="#1f4e78", bd=1, relief=tk.RAISED)
        top_frame.pack(fill=tk.X, side=tk.TOP, padx=8, pady=5)
        
        tk.Label(top_frame, text="Folder Foto:", bg="#1f4e78", fg="white", font=("Arial", 10, "bold")).pack(side=tk.LEFT, padx=5, pady=5)
        self.ent_path = tk.Entry(top_frame, width=45, font=("Arial", 10))
        self.ent_path.pack(side=tk.LEFT, padx=5, pady=5, fill=tk.X, expand=True)
        
        tk.Button(top_frame, text="PILIH", bg="#3498db", fg="white", font=("Arial", 9, "bold"), width=8, command=self.pilih_folder).pack(side=tk.LEFT, padx=3, pady=5)
        tk.Button(top_frame, text="MUAT", bg="#2ecc71", fg="white", font=("Arial", 9, "bold"), width=8, command=self.muat_foto).pack(side=tk.LEFT, padx=3, pady=5)

        self.ai_mode_button = tk.Button(
            top_frame, text="MODE AI: OFF", bg="#34495e", fg="white", font=("Arial", 9, "bold"), command=self.toggle_ai_mode
        )
        self.ai_mode_button.pack(side=tk.RIGHT, padx=5, pady=5)

        tk.Button(
            top_frame, text="PANDUAN", bg="#16a085", fg="white", font=("Arial", 9, "bold"), command=self.show_help_and_diagnostics
        ).pack(side=tk.RIGHT, padx=5, pady=5)
        
        tk.Button(top_frame, text="PREVIEW", bg="#f1c40f", fg="black", font=("Arial", 9, "bold"), width=9, command=self.tampilkan_preview_popup).pack(side=tk.RIGHT, padx=5, pady=5)
        tk.Button(top_frame, text="EKSPORT", bg="#e67e22", fg="white", font=("Arial", 9, "bold"), width=9, command=self.eksport_ke_excel).pack(side=tk.RIGHT, padx=5, pady=5)
        
        # Main Layout (3 Kolom Sejajar + Panel Bawah)
        main_layout = tk.Frame(self.survey_tab, bg="#eaf2f8")
        main_layout.pack(fill=tk.BOTH, expand=True, padx=8, pady=5)
        
        row_top_panels = tk.Frame(main_layout, bg="#eaf2f8")
        row_top_panels.pack(fill=tk.BOTH, expand=True)
        row_top_panels.columnconfigure(0, weight=2)
        row_top_panels.columnconfigure(1, weight=12)
        row_top_panels.columnconfigure(2, weight=2)
        row_top_panels.rowconfigure(0, weight=1)

        def make_scroll_panel(parent, column):
            outer = tk.Frame(parent, bg="#eaf2f8", bd=1, relief=tk.GROOVE)
            outer.grid(row=0, column=column, sticky="nsew", padx=3, pady=3)
            
            canvas = tk.Canvas(outer, borderwidth=0, background="#f5f6fa", highlightthickness=0)
            scrollbar = tk.Scrollbar(outer, orient="vertical", command=canvas.yview)
            content = tk.Frame(canvas, bg="#f5f6fa")
            
            content.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
            window_id = canvas.create_window((0, 0), window=content, anchor="nw")
            canvas.bind("<Configure>", lambda e: canvas.itemconfigure(window_id, width=e.width))
            canvas.configure(yscrollcommand=scrollbar.set)
            
            canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
            scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
            return content

        # Keyboard Bindings
        self.root.bind("<Control-r>", lambda event: self.reset_penilaian())
        self.root.bind("<Control-R>", lambda event: self.reset_penilaian())
        self.root.bind("<Left>", lambda event: self.foto_sebelumnya())
        self.root.bind("<Right>", lambda event: self.simpan_dan_selanjutnya())

        # Kolom Kiri (Bahu & Drainase Kiri)
        left_column = make_scroll_panel(row_top_panels, 0)
        tk.Label(left_column, text="KIRI - BAHU JALAN", bg="#7f8c8d", fg="white", font=("Arial", 10, "bold")).pack(fill=tk.X, padx=2, pady=4)
        
        f_lebar_bk = tk.Frame(left_column, bg="#f5f6fa")
        f_lebar_bk.pack(fill=tk.X, padx=5, pady=2)
        tk.Label(f_lebar_bk, text="LEBAR (m):", bg="#f5f6fa", font=("Arial", 9, "bold")).pack(side=tk.LEFT)
        tk.Entry(f_lebar_bk, textvariable=self.var_lebar_bahu_kiri, width=12, font=("Arial", 10)).pack(side=tk.LEFT, padx=5)
        
        f_matrix_bk = tk.Frame(left_column, bg="#f5f6fa")
        f_matrix_bk.pack(fill=tk.X, padx=2, pady=2)
        f_matrix_bk.columnconfigure(0, weight=1)
        f_matrix_bk.columnconfigure(1, weight=1)
        self.create_side_matrix(f_matrix_bk, "TIPE", "bahu_kiri_tipe", 0)
        self.create_side_matrix(f_matrix_bk, "KONDISI", "bahu_kiri_kondisi", 1)
        
        tk.Label(left_column, text="KIRI - DRAINASE", bg="#7f8c8d", fg="white", font=("Arial", 10, "bold")).pack(fill=tk.X, padx=2, pady=(12, 4))
        f_lebar_dk = tk.Frame(left_column, bg="#f5f6fa")
        f_lebar_dk.pack(fill=tk.X, padx=5, pady=2)
        tk.Label(f_lebar_dk, text="LEBAR (m):", bg="#f5f6fa", font=("Arial", 9, "bold")).pack(side=tk.LEFT)
        tk.Entry(f_lebar_dk, textvariable=self.var_lebar_drainase_kiri, width=12, font=("Arial", 10)).pack(side=tk.LEFT, padx=5)
        
        f_matrix_dk = tk.Frame(left_column, bg="#f5f6fa")
        f_matrix_dk.pack(fill=tk.X, padx=2, pady=2)
        f_matrix_dk.columnconfigure(0, weight=1)
        f_matrix_dk.columnconfigure(1, weight=1)
        self.create_side_matrix(f_matrix_dk, "TIPE", "drainase_kiri_tipe", 0)
        self.create_side_matrix(f_matrix_dk, "KONDISI", "drainase_kiri_kondisi", 1)

        # Kolom Tengah (Canvas Foto)
        center_column = tk.Frame(row_top_panels, bg="#ffffff", bd=1, relief=tk.SOLID)
        center_column.grid(row=0, column=1, sticky="nsew", padx=3, pady=3)
        center_column.rowconfigure(0, weight=0)
        center_column.rowconfigure(1, weight=1)
        center_column.columnconfigure(0, weight=1)
        
        photo_toolbar = tk.Frame(center_column, bg="#34495e")
        photo_toolbar.grid(row=0, column=0, sticky="ew")
        
        self.lbl_photo_info = tk.Label(photo_toolbar, text="Foto: 0 / 0 | -", bg="#34495e", fg="white", font=("Arial", 9, "bold"))
        self.lbl_photo_info.pack(side=tk.LEFT, padx=10, pady=5)

        self.lbl_ai_status = tk.Label(photo_toolbar, textvariable=self.ai_status, bg="#34495e", fg="#f1c40f", font=("Arial", 9, "bold"))
        self.lbl_ai_status.pack(side=tk.LEFT, padx=8, pady=5)
        
        tk.Button(photo_toolbar, text=" ⛶ ", font=("Arial", 9, "bold"), bg="#95a5a6", fg="white", width=4, command=self.zoom_fit).pack(side=tk.RIGHT, padx=4, pady=5)
        tk.Button(photo_toolbar, text=" - ", font=("Arial", 9, "bold"), bg="#e74c3c", fg="white", width=4, command=self.zoom_out).pack(side=tk.RIGHT, padx=4, pady=5)
        tk.Button(photo_toolbar, text=" + ", font=("Arial", 9, "bold"), bg="#2ecc71", fg="white", width=4, command=self.zoom_in).pack(side=tk.RIGHT, padx=4, pady=5)

        canvas_container = tk.Frame(center_column, bg="grey")
        canvas_container.grid(row=1, column=0, sticky="nsew")
        canvas_container.rowconfigure(0, weight=1)
        canvas_container.columnconfigure(0, weight=1)

        self.canvas = tk.Canvas(canvas_container, bg="#2f3640", highlightthickness=0)
        self.canvas.grid(row=0, column=0, sticky="nsew")
        
        h_scroll = tk.Scrollbar(canvas_container, orient=tk.HORIZONTAL, command=self.canvas.xview)
        h_scroll.grid(row=1, column=0, sticky="ew")
        v_scroll = tk.Scrollbar(canvas_container, orient=tk.VERTICAL, command=self.canvas.yview)
        v_scroll.grid(row=0, column=1, sticky="ns")
        self.canvas.configure(xscrollcommand=h_scroll.set, yscrollcommand=v_scroll.set)

        self.canvas.bind("<ButtonPress-1>", self.start_pan)
        self.canvas.bind("<B1-Motion>", self.execute_pan)
        self.canvas.bind("<MouseWheel>", self.zoom_pake_scroll_mouse)
        self.canvas.bind("<Button-4>", self.zoom_pake_scroll_mouse)
        self.canvas.bind("<Button-5>", self.zoom_pake_scroll_mouse)

        # Kolom Kanan (Bahu & Drainase Kanan)
        right_column = make_scroll_panel(row_top_panels, 2)
        tk.Label(right_column, text="KANAN - BAHU JALAN", bg="#7f8c8d", fg="white", font=("Arial", 10, "bold")).pack(fill=tk.X, padx=2, pady=4)

        f_lebar_bkr = tk.Frame(right_column, bg="#f5f6fa")
        f_lebar_bkr.pack(fill=tk.X, padx=5, pady=2)
        tk.Label(f_lebar_bkr, text="LEBAR (m):", bg="#f5f6fa", font=("Arial", 9, "bold")).pack(side=tk.LEFT)
        tk.Entry(f_lebar_bkr, textvariable=self.var_lebar_bahu_kanan, width=12, font=("Arial", 10)).pack(side=tk.LEFT, padx=5)

        f_matrix_bk_r = tk.Frame(right_column, bg="#f5f6fa")
        f_matrix_bk_r.pack(fill=tk.X, padx=2, pady=2)
        f_matrix_bk_r.columnconfigure(0, weight=1)
        f_matrix_bk_r.columnconfigure(1, weight=1)
        self.create_side_matrix(f_matrix_bk_r, "TIPE", "bahu_kanan_tipe", 0)
        self.create_side_matrix(f_matrix_bk_r, "KONDISI", "bahu_kanan_kondisi", 1)

        tk.Label(right_column, text="KANAN - DRAINASE", bg="#7f8c8d", fg="white", font=("Arial", 10, "bold")).pack(fill=tk.X, padx=2, pady=(12, 4))
        f_lebar_dkr = tk.Frame(right_column, bg="#f5f6fa")
        f_lebar_dkr.pack(fill=tk.X, padx=5, pady=2)
        tk.Label(f_lebar_dkr, text="LEBAR (m):", bg="#f5f6fa", font=("Arial", 9, "bold")).pack(side=tk.LEFT)
        tk.Entry(f_lebar_dkr, textvariable=self.var_lebar_drainase_kanan, width=12, font=("Arial", 10)).pack(side=tk.LEFT, padx=5)

        f_matrix_dk_r = tk.Frame(right_column, bg="#f5f6fa")
        f_matrix_dk_r.pack(fill=tk.X, padx=2, pady=2)
        f_matrix_dk_r.columnconfigure(0, weight=1)
        f_matrix_dk_r.columnconfigure(1, weight=1)
        self.create_side_matrix(f_matrix_dk_r, "TIPE", "drainase_kanan_tipe", 0)
        self.create_side_matrix(f_matrix_dk_r, "KONDISI", "drainase_kanan_kondisi", 1)

        # Bottom Control Master Layout
        row_bottom_master = tk.Frame(main_layout, bg="#eaf2f8")
        row_bottom_master.pack(fill=tk.X, side=tk.BOTTOM, pady=(10, 5))
        row_bottom_master.columnconfigure(0, weight=1)
        row_bottom_master.columnconfigure(1, weight=8)
        row_bottom_master.columnconfigure(2, weight=1)

        # Tombol Kiri (Reset & Lewati)
        frame_buttons_left = tk.Frame(row_bottom_master, bg="#eaf2f8")
        frame_buttons_left.grid(row=0, column=0, sticky="nsew", padx=(5, 10))
        tk.Button(frame_buttons_left, text="RESET", bg="#e74c3c", fg="white", font=("Arial", 9, "bold"), command=self.reset_penilaian).pack(fill=tk.BOTH, expand=True, pady=(0, 5))
        tk.Button(frame_buttons_left, text="LEWATI", bg="#95a5a6", fg="white", font=("Arial", 9, "bold"), command=self.lewati_tanpa_simpan).pack(fill=tk.BOTH, expand=True)

        # Perkerasan Utama (Tengah Bawah)
        row_bottom_perkerasan = tk.Frame(row_bottom_master, bg="#eaf2f8", bd=1, relief=tk.GROOVE, pady=1)
        row_bottom_perkerasan.grid(row=0, column=1, sticky="nsew")

        tk.Label(row_bottom_perkerasan, text="PERKERASAN UTAMA", bg="#1976d2", fg="white", font=("Arial", 10, "bold")).pack(fill=tk.X, padx=6, pady=2)
        
        f_lebar_p = tk.Frame(row_bottom_perkerasan, bg="#eaf2f8")
        f_lebar_p.pack(fill=tk.X, padx=5, pady=2)
        tk.Label(f_lebar_p, text="LEBAR PERKERASAN (m):", bg="#eaf2f8", font=("Arial", 9, "bold")).pack(side=tk.LEFT)
        tk.Entry(f_lebar_p, textvariable=self.var_lebar_perkerasan, width=15, font=("Arial", 10)).pack(side=tk.LEFT, padx=5)

        self.create_perkerasan_buttons(row_bottom_perkerasan, "TIPE PERKERASAN", "perkerasan_tipe")
        self.create_perkerasan_buttons(row_bottom_perkerasan, "KONDISI PERKERASAN", "perkerasan_kondisi")

        # Tombol Kanan (Sebelumnya & Selanjutnya)
        frame_buttons_right = tk.Frame(row_bottom_master, bg="#eaf2f8")
        frame_buttons_right.grid(row=0, column=2, sticky="nsew", padx=(10, 5))
        tk.Button(frame_buttons_right, text="SEBELUMNYA", bg="#7f8c8d", fg="white", font=("Arial", 9, "bold"), command=self.foto_sebelumnya).pack(fill=tk.BOTH, expand=True, pady=(0, 5))
        tk.Button(frame_buttons_right, text="SIMPAN & LANJUT", bg="#27ae60", fg="white", font=("Arial", 9, "bold"), command=self.simpan_dan_selanjutnya).pack(fill=tk.BOTH, expand=True)

        self.build_ai_tab()

    def build_pkrms_tab(self):
        """Antarmuka PKRMS: template, audit dataset, dan ekspor catatan survei."""
        self.pkrms_tab.columnconfigure(0, weight=1)
        ttk.Label(
            self.pkrms_tab,
            text="PKRMS — Pengelolaan Data Survei Jalan",
            font=("Segoe UI", 15, "bold"),
        ).grid(row=0, column=0, sticky="w", pady=(0, 4))
        ttk.Label(
            self.pkrms_tab,
            text="Siapkan template, audit foto/dataset, dan ekspor catatan survei ke Excel. "
                 "Data yang belum dinilai tidak diberi label kondisi secara otomatis.",
            wraplength=1000,
        ).grid(row=1, column=0, sticky="w", pady=(0, 12))
        actions = ttk.LabelFrame(self.pkrms_tab, text="Operasi PKRMS")
        actions.grid(row=2, column=0, sticky="new", pady=5)
        ttk.Button(actions, text="Buat Template PKRMS Excel", command=self.pkrms_create_template).grid(row=0, column=0, padx=6, pady=8, sticky="w")
        ttk.Button(actions, text="Ekspor Data Survei ke PKRMS", command=self.pkrms_export_report).grid(row=0, column=1, padx=6, pady=8, sticky="w")
        ttk.Button(actions, text="Audit Folder Foto", command=self.pkrms_audit_dataset).grid(row=0, column=2, padx=6, pady=8, sticky="w")
        ttk.Button(actions, text="Audit Dataset Klasifikasi", command=self.pkrms_audit_classified_dataset).grid(row=1, column=0, padx=6, pady=8, sticky="w")
        ttk.Button(actions, text="Ekspor Rekap Dataset Klasifikasi", command=self.pkrms_export_classified_dataset).grid(row=1, column=1, padx=6, pady=8, sticky="w")
        self.pkrms_status = tk.StringVar(value="Siap. Pilih operasi PKRMS yang diperlukan.")
        ttk.Label(self.pkrms_tab, textvariable=self.pkrms_status, wraplength=1000, justify="left").grid(row=3, column=0, sticky="nw", pady=10)
        ttk.Label(
            self.pkrms_tab,
            text="Catatan: template ini adalah struktur kerja dan bukan pengganti validasi terhadap format resmi PKRMS.",
            wraplength=1000,
        ).grid(row=4, column=0, sticky="w", pady=(8, 0))

    def pkrms_create_template(self):
        destination = filedialog.asksaveasfilename(
            parent=self.root, title="Simpan Template PKRMS", defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")], initialfile="template_pkrms.xlsx",
        )
        if not destination:
            return
        try:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "PKRMS_Survey"
            columns = [
                "province", "regency", "district", "road_code", "road_name",
                "segment_id", "start_km", "end_km", "length_km", "survey_date",
                "photo_name", "surface_type", "condition", "ai_condition",
                "corrected_condition", "final_condition", "review_status",
                "road_width_m", "width_measurement_status", "remarks",
            ]
            sheet.append(columns)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = "A1:T1"
            for cell in sheet[1]:
                cell.font = XlFont(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = XlAlignment(wrap_text=True, vertical="center")
            for idx, name in enumerate(columns, start=1):
                sheet.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = max(14, min(26, len(name) + 3))
            workbook.save(destination)
            self.pkrms_status.set(f"Template PKRMS berhasil dibuat: {destination}")
            messagebox.showinfo("PKRMS", "Template Excel berhasil dibuat.", parent=self.root)
        except Exception as error:
            log_application_error("Membuat template PKRMS", type(error), error, error.__traceback__)
            messagebox.showerror("PKRMS", f"Gagal membuat template:\n{error}", parent=self.root)

    def pkrms_export_report(self):
        destination = filedialog.asksaveasfilename(
            parent=self.root, title="Ekspor Data Survei untuk PKRMS", defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")], initialfile="data_survei_pkrms.xlsx",
        )
        if not destination:
            return
        try:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "PKRMS_Survey"
            sheet.append(["photo_name", "photo_path", "survey_date", "surface_type", "condition", "review_status", "remarks"])
            records = dict(self._stored_survey_records)
            records.update(self.survey_data)
            for photo_key, data in records.items():
                data = data if isinstance(data, dict) else {}
                sheet.append([
                    data.get("filename", os.path.basename(str(photo_key))),
                    data.get("folder", os.path.dirname(str(photo_key))),
                    data.get("survey_date", ""),
                    data.get("jenis_perkerasan", data.get("perkerasan", "")),
                    data.get("kondisi", data.get("kondisi_jalan", "")),
                    data.get("review_status", "Belum diverifikasi"),
                    data.get("remarks", ""),
                ])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = XlFont(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for col in sheet.columns:
                letter = col[0].column_letter
                sheet.column_dimensions[letter].width = min(45, max(14, max(len(str(cell.value or "")) for cell in col) + 2))
            workbook.save(destination)
            self.pkrms_status.set(f"Ekspor selesai: {len(records)} catatan ditulis ke {destination}")
            messagebox.showinfo("PKRMS", f"Ekspor selesai. Jumlah catatan: {len(records)}.", parent=self.root)
        except Exception as error:
            log_application_error("Ekspor data PKRMS", type(error), error, error.__traceback__)
            messagebox.showerror("PKRMS", f"Gagal mengekspor data:\n{error}", parent=self.root)

    def pkrms_audit_classified_dataset(self):
        """Gunakan utilitas dataset klasifikasi dari implementasi modular lama."""
        try:
            report = pkrms_utils.audit_dataset()
            lines = ["Audit folder dataset klasifikasi (APP_ROOT/dataset):", ""]
            lines.append("Kondisi jalan:")
            lines.extend(f"  {label}: {count} foto" for label, count in report["condition_counts"].items())
            lines.append("")
            lines.append("Jenis perkerasan:")
            lines.extend(f"  {label}: {count} foto" for label, count in report["surface_counts"].items())
            lines.extend([
                "",
                "Audit ini menghitung file gambar berdasarkan folder kelas.",
                "Audit tidak membuktikan label benar atau model AI telah dilatih.",
            ])
            result = "\n".join(lines)
            self.pkrms_status.set(result)
            messagebox.showinfo("Audit Dataset Klasifikasi", result, parent=self.root)
        except Exception as error:
            log_application_error("Audit dataset klasifikasi", type(error), error, error.__traceback__)
            messagebox.showerror("Audit Dataset", str(error), parent=self.root)

    def pkrms_export_classified_dataset(self):
        destination = filedialog.asksaveasfilename(
            parent=self.root,
            title="Ekspor Rekap Dataset Klasifikasi",
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")],
            initialfile="rekap_dataset_klasifikasi.xlsx",
        )
        if not destination:
            return
        try:
            counts = pkrms_utils.export_pkrms_report(destination)
            summary = ", ".join(f"{label}: {count}" for label, count in counts.items())
            message = f"Rekap dataset diekspor ke:\n{destination}\n\nJumlah foto per kelas: {summary}"
            self.pkrms_status.set(message)
            messagebox.showinfo("Ekspor Dataset", message, parent=self.root)
        except Exception as error:
            log_application_error("Ekspor rekap dataset klasifikasi", type(error), error, error.__traceback__)
            messagebox.showerror("Ekspor Dataset", str(error), parent=self.root)

    def pkrms_audit_dataset(self):
        try:
            extensions = {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".tif", ".tiff"}
            image_count = sum(
                1 for path in SURVEY_AI_IMAGES.rglob("*")
                if path.is_file() and path.suffix.lower() in extensions
            ) if SURVEY_AI_IMAGES.is_dir() else 0
            manifest_count = len(self.training_records)
            saved_count = len(self._stored_survey_records)
            report = (
                "Audit dataset selesai.\n\n"
                f"Foto di dataset AI: {image_count}\n"
                f"Rekaman label/manifest AI: {manifest_count}\n"
                f"Catatan survei tersimpan: {saved_count}\n\n"
                "Audit ini menghitung berkas dan rekaman yang ditemukan; belum memverifikasi "
                "kebenaran label atau kesesuaian dengan skema resmi PKRMS."
            )
            self.pkrms_status.set(report)
            messagebox.showinfo("Audit Dataset PKRMS", report, parent=self.root)
        except Exception as error:
            log_application_error("Audit dataset PKRMS", type(error), error, error.__traceback__)
            messagebox.showerror("PKRMS", f"Audit gagal:\n{error}", parent=self.root)

    def build_online_tab(self):
        """Integrasi cloud/Colab dan sinkronisasi PostgreSQL dari frontend lama."""
        self.online_tab.columnconfigure(0, weight=1)
        ttk.Label(
            self.online_tab,
            text="Integrasi Online, Google Colab, dan Database",
            font=("Segoe UI", 15, "bold"),
        ).grid(row=0, column=0, sticky="w", pady=(0, 4))
        ttk.Label(
            self.online_tab,
            text="Fitur ini membutuhkan koneksi internet dan layanan eksternal. "
                 "Aplikasi survei lokal tetap dapat digunakan secara offline.",
            wraplength=1000,
        ).grid(row=1, column=0, sticky="w", pady=(0, 12))

        colab_frame = ttk.LabelFrame(self.online_tab, text="Google Colab")
        colab_frame.grid(row=2, column=0, sticky="ew", pady=6)
        ttk.Label(
            colab_frame,
            text="Buat skrip awal untuk menyiapkan Google Drive dan lingkungan Colab. "
                 "Skrip ini bukan pipeline training AI lengkap.",
            wraplength=950,
        ).grid(row=0, column=0, columnspan=2, sticky="w", padx=8, pady=(8, 4))
        ttk.Button(
            colab_frame,
            text="Ekspor Skrip Awal Colab (.py)",
            command=self.export_colab_script,
        ).grid(row=1, column=0, sticky="w", padx=8, pady=8)

        db_frame = ttk.LabelFrame(self.online_tab, text="Sinkronisasi PostgreSQL")
        db_frame.grid(row=3, column=0, sticky="ew", pady=6)
        fields = [
            ("Host", "db_host", "localhost"),
            ("Port", "db_port", "5432"),
            ("Database", "db_name", "aprj_db"),
            ("User", "db_user", "postgres"),
            ("Password", "db_password", ""),
        ]
        self.db_vars = {}
        for index, (label, key, default) in enumerate(fields):
            row, col = divmod(index, 4)
            ttk.Label(db_frame, text=label).grid(row=row * 2, column=col, sticky="w", padx=8, pady=(8, 2))
            variable = tk.StringVar(value=default)
            self.db_vars[key] = variable
            entry = ttk.Entry(
                db_frame, textvariable=variable, width=22,
                show="*" if key == "db_password" else "",
            )
            entry.grid(row=row * 2 + 1, column=col, sticky="ew", padx=8, pady=(0, 8))
        for col in range(4):
            db_frame.columnconfigure(col, weight=1)
        controls = ttk.Frame(db_frame)
        controls.grid(row=4, column=0, columnspan=4, sticky="w", padx=8, pady=(0, 8))
        ttk.Button(controls, text="Uji Koneksi", command=self.test_postgres_connection).pack(side="left", padx=(0, 6))
        ttk.Button(controls, text="Sinkronkan Data Survei", command=self.sync_postgres).pack(side="left")

        self.online_status = tk.StringVar(value="Belum ada operasi online yang dijalankan.")
        ttk.Label(
            self.online_tab, textvariable=self.online_status,
            wraplength=1000, justify="left",
        ).grid(row=4, column=0, sticky="w", pady=8)

    def _make_database_manager(self):
        try:
            port = int(self.db_vars["db_port"].get().strip())
            if not 1 <= port <= 65535:
                raise ValueError
        except ValueError:
            raise ValueError("Port PostgreSQL harus berupa angka antara 1 dan 65535.")
        return database.DatabaseManager(
            host=self.db_vars["db_host"].get().strip() or "localhost",
            port=port,
            database=self.db_vars["db_name"].get().strip() or "aprj_db",
            user=self.db_vars["db_user"].get().strip() or "postgres",
            password=self.db_vars["db_password"].get(),
        )

    def test_postgres_connection(self):
        try:
            manager = self._make_database_manager()
            success, message = manager.test_connection()
            self.online_status.set(message)
            if success:
                messagebox.showinfo("PostgreSQL", message, parent=self.root)
            else:
                messagebox.showerror("PostgreSQL", message, parent=self.root)
        except Exception as error:
            messagebox.showerror("Konfigurasi PostgreSQL", str(error), parent=self.root)

    def sync_postgres(self):
        """Sinkronkan catatan yang sudah tersimpan; tidak mengirim input yang belum disimpan."""
        try:
            records = dict(self._stored_survey_records)
            records.update(self.survey_data)
            normalized = {}
            for key, value in records.items():
                if not isinstance(value, dict):
                    continue
                data = value.copy()
                source_path = str(key)
                data.setdefault("filename", os.path.basename(source_path))
                data.setdefault("folder", os.path.dirname(source_path) or self.photo_folder)
                normalized[source_path] = data
            if not normalized:
                messagebox.showinfo(
                    "PostgreSQL", "Belum ada catatan survei tersimpan untuk disinkronkan.",
                    parent=self.root,
                )
                return
            manager = self._make_database_manager()
            success, message = manager.sync_to_postgres(normalized)
            self.online_status.set(message)
            if success:
                messagebox.showinfo(
                    "PostgreSQL", f"{message}\nJumlah catatan: {len(normalized)}",
                    parent=self.root,
                )
            else:
                messagebox.showerror("PostgreSQL", message, parent=self.root)
        except Exception as error:
            log_application_error("Sinkronisasi PostgreSQL", type(error), error, error.__traceback__)
            messagebox.showerror("PostgreSQL", str(error), parent=self.root)

    def export_colab_script(self):
        destination = filedialog.asksaveasfilename(
            parent=self.root,
            title="Ekspor Skrip Awal Google Colab",
            defaultextension=".py",
            initialfile="train_colab_aprj.py",
            filetypes=[("Python Files", "*.py")],
        )
        if not destination:
            return
        try:
            success, message = ai_model.GoogleDriveColabManager.generate_colab_training_notebook(destination)
            if success:
                self.online_status.set(f"{message} File: {destination}")
                messagebox.showinfo(
                    "Google Colab",
                    f"{message}\n\nCatatan: skrip yang dihasilkan adalah kerangka awal, "
                    "belum melakukan training model end-to-end.\n\nDisimpan di:\n{destination}",
                    parent=self.root,
                )
            else:
                messagebox.showerror("Google Colab", message, parent=self.root)
        except Exception as error:
            log_application_error("Ekspor skrip Colab", type(error), error, error.__traceback__)
            messagebox.showerror("Google Colab", str(error), parent=self.root)

    def load_training_manifest(self):
        if not SURVEY_AI_MANIFEST.exists():
            return []
        try:
            data = json.loads(SURVEY_AI_MANIFEST.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return []
        return [record for record in data if isinstance(record, dict)]

    def save_training_manifest(self):
        SURVEY_AI_DATASET.mkdir(parents=True, exist_ok=True)
        temporary = SURVEY_AI_MANIFEST.with_suffix(".json.tmp")
        temporary.write_text(json.dumps(self.training_records, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(SURVEY_AI_MANIFEST)

    def build_ai_tab(self):
        self.ai_tab.columnconfigure(0, weight=1)
        self.ai_tab.rowconfigure(2, weight=1)
        ttk.Label(self.ai_tab, text="Rekayasa & Pelatihan AI Survei Jalan", font=("Segoe UI", 15, "bold")).grid(row=0, column=0, sticky="w", pady=(0, 4))
        ttk.Label(self.ai_tab, text="Kelola dataset terkonfirmasi manusia dan latih model multi-task AI secara lokal.").grid(row=1, column=0, sticky="w", pady=(0, 10))

        dataset_frame = ttk.LabelFrame(self.ai_tab, text="Dataset Terkonfirmasi")
        dataset_frame.grid(row=2, column=0, sticky="nsew", pady=5)
        dataset_frame.columnconfigure(0, weight=1)
        dataset_frame.rowconfigure(0, weight=1)
        
        columns = ("photo", "source", "categories", "widths", "state")
        self.ai_dataset_tree = ttk.Treeview(dataset_frame, columns=columns, show="headings", height=10)
        for key, heading, width in (
            ("photo", "Foto", 200),
            ("source", "Path Dataset", 400),
            ("categories", "Kategori", 100),
            ("widths", "Lebar", 100),
            ("state", "Status", 150),
        ):
            self.ai_dataset_tree.heading(key, text=heading)
            self.ai_dataset_tree.column(key, width=width, anchor="w")
        self.ai_dataset_tree.grid(row=0, column=0, sticky="nsew")
        
        tree_scroll = ttk.Scrollbar(dataset_frame, orient="vertical", command=self.ai_dataset_tree.yview)
        tree_scroll.grid(row=0, column=1, sticky="ns")
        self.ai_dataset_tree.configure(yscrollcommand=tree_scroll.set)

        import_controls = ttk.Frame(dataset_frame)
        import_controls.grid(row=1, column=0, columnspan=2, sticky="ew", pady=5)
        ttk.Button(import_controls, text="Pilih Export Excel", command=self.stage_survey_exports).pack(side="left", padx=3)
        self.confirm_import_button = ttk.Button(import_controls, text="Konfirmasi Impor", command=self.confirm_staged_import, state="disabled")
        self.confirm_import_button.pack(side="left", padx=3)
        self.cancel_import_button = ttk.Button(import_controls, text="Batalkan", command=self.cancel_staged_import, state="disabled")
        self.cancel_import_button.pack(side="left", padx=3)
        ttk.Button(import_controls, text="Validasi Dataset", command=self.validate_training_dataset).pack(side="left", padx=3)

        training_frame = ttk.LabelFrame(self.ai_tab, text="Kontrol Training AI")
        training_frame.grid(row=3, column=0, sticky="ew", pady=5)
        
        ttk.Label(training_frame, text="Device:").grid(row=0, column=0, padx=5, pady=6)
        self.device_combo = ttk.Combobox(training_frame, textvariable=self.ai_device_var, values=("Otomatis", "CPU", "GPU (CUDA)"), state="readonly", width=12)
        self.device_combo.grid(row=0, column=1, padx=4)
        
        ttk.Label(training_frame, text="Batch:").grid(row=0, column=2, padx=5)
        ttk.Entry(training_frame, textvariable=self.ai_batch_var, width=5).grid(row=0, column=3)
        
        ttk.Label(training_frame, text="Epoch:").grid(row=0, column=4, padx=5)
        ttk.Entry(training_frame, textvariable=self.ai_epochs_var, width=5).grid(row=0, column=5)
        
        ttk.Button(training_frame, text="Mulai Training", command=self.start_survey_training).grid(row=0, column=6, padx=5, pady=5)
        
        self.cancel_training_button = ttk.Button(training_frame, text="Batalkan", command=self.cancel_survey_training, state="disabled")
        self.cancel_training_button.grid(row=0, column=7, padx=5)
        
        ttk.Button(training_frame, text="Buka Folder AI", command=self.open_ai_folder).grid(row=0, column=8, padx=5)
        ttk.Button(training_frame, text="Periksa Kesiapan Model", command=self.check_ai_inference_readiness).grid(row=1, column=0, columnspan=3, padx=5, pady=5, sticky="w")
        
        ttk.Label(training_frame, textvariable=self.ai_training_status, wraplength=900).grid(row=2, column=0, columnspan=9, sticky="w", padx=5, pady=5)
        self.refresh_training_dataset_view()

    def refresh_training_dataset_view(self):
        if not hasattr(self, "ai_dataset_tree"):
            return
        self.ai_dataset_tree.delete(*self.ai_dataset_tree.get_children())
        for record in self.training_records:
            image_path = APP_ROOT / record.get("image_path", "")
            labels = record.get("labels", {})
            category_count = sum(bool(labels.get(task)) for task in SURVEY_TASKS)
            width_count = sum(labels.get(field) is not None for field in SURVEY_WIDTH_FIELDS)
            self.ai_dataset_tree.insert(
                "", tk.END,
                values=(record.get("photo_name", image_path.name), str(image_path), f"{category_count}/10", f"{width_count}/5", "Tersedia" if image_path.is_file() else "Hilang")
            )

    def stage_survey_exports(self):
        paths = filedialog.askopenfilenames(title="Pilih file Excel hasil survei", filetypes=[("Excel", "*.xlsx")])
        if not paths:
            return
        staged = []
        try:
            for workbook_path in paths:
                wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
                if "Hasil Survei Jalan" not in wb.sheetnames:
                    continue
                sheet = wb["Hasil Survei Jalan"]
                for row_idx, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
                    if len(row) < 17 or not row[1]:
                        continue
                    labels = {
                        "bahu_kiri_tipe": str(row[3] or "").strip() or None,
                        "bahu_kiri_kondisi": str(row[4] or "").strip() or None,
                        "drainase_kiri_tipe": str(row[6] or "").strip() or None,
                        "drainase_kiri_kondisi": str(row[7] or "").strip() or None,
                        "perkerasan_tipe": str(row[9] or "").strip() or None,
                        "perkerasan_kondisi": str(row[10] or "").strip() or None,
                        "bahu_kanan_tipe": str(row[12] or "").strip() or None,
                        "bahu_kanan_kondisi": str(row[13] or "").strip() or None,
                        "drainase_kanan_tipe": str(row[15] or "").strip() or None,
                        "drainase_kanan_kondisi": str(row[16] or "").strip() or None,
                    }
                    source_folder = Path(str(row[0])) if row[0] else Path()
                    photo_relative_path = str(row[1]).strip()
                    photo_name = Path(photo_relative_path).name
                    photo_path = source_folder / photo_relative_path
                    if photo_path.is_file():
                        staged.append({
                            "photo_name": photo_name,
                            "source_path": str(photo_path),
                            "workbook": str(workbook_path),
                            "excel_row": row_idx,
                            "labels": labels,
                        })
                wb.close()
        except Exception as e:
            messagebox.showerror("Gagal Impor Excel", str(e))
            return

        if not staged:
            messagebox.showwarning("Peringatan", "Tidak ditemukan data valid dalam file Excel.")
            return

        self.staged_training_records = staged
        self.confirm_import_button.configure(state="normal")
        self.cancel_import_button.configure(state="normal")
        self.ai_training_status.set(f"Pratinjau: {len(staged)} foto siap diimpor. Klik Konfirmasi.")

    def cancel_staged_import(self):
        self.staged_training_records = []
        self.confirm_import_button.configure(state="disabled")
        self.cancel_import_button.configure(state="disabled")
        self.ai_training_status.set("Impor dibatalkan.")

    def confirm_staged_import(self):
        if not self.staged_training_records:
            return
        added = 0
        try:
            for staged in self.staged_training_records:
                source = Path(staged["source_path"])
                digest = hashlib.sha256(source.read_bytes()).hexdigest()
                destination = SURVEY_AI_IMAGES / f"{digest[:20]}_{source.name}"
                if not destination.exists():
                    shutil.copy2(source, destination)
                record = {
                    "photo_name": staged["photo_name"],
                    "image_path": str(destination.relative_to(APP_ROOT)),
                    "sha256": digest,
                    "labels": staged["labels"],
                    "imported_at": datetime.now().isoformat(),
                }
                self.training_records.append(record)
                added += 1
            self.save_training_manifest()
        except Exception as e:
            messagebox.showerror("Gagal Menyalin", str(e))
            return

        self.staged_training_records = []
        self.confirm_import_button.configure(state="disabled")
        self.cancel_import_button.configure(state="disabled")
        self.refresh_training_dataset_view()
        self.ai_training_status.set(f"Berhasil mengimpor {added} foto ke dataset AI.")

    def validate_training_dataset(self):
        available = sum(1 for r in self.training_records if (APP_ROOT / r.get("image_path", "")).is_file())
        messagebox.showinfo("Validasi Dataset", f"Total record terdaftar: {len(self.training_records)}\nFile foto fisik tersedia: {available}")

    def resolve_training_device(self, torch_module):
        choice = self.ai_device_var.get()
        if choice == "CPU":
            return torch_module.device("cpu")
        if choice == "GPU (CUDA)":
            if not torch_module.cuda.is_available():
                raise RuntimeError("CUDA tidak tersedia di perangkat ini.")
            return torch_module.device("cuda")
        return torch_module.device("cuda" if torch_module.cuda.is_available() else "cpu")

    def start_survey_training(self):
        if self.training_running:
            return
        messagebox.showinfo("Training AI belum tersedia", "Antarmuka pengelolaan dataset tersedia, tetapi pipeline training belum diimplementasikan pada file ini. Jangan anggap model AI sudah terlatih.")

    def check_ai_inference_readiness(self):
        """Laporkan kesiapan inferensi tanpa menghasilkan prediksi buatan."""
        model_path = SURVEY_AI_MODEL
        if not ai_model.TORCH_AVAILABLE:
            message = "PyTorch belum tersedia. Inferensi tidak dapat dijalankan."
        elif not model_path.is_file():
            message = (
                "PyTorch tersedia, tetapi file model terlatih belum ditemukan:\n"
                f"{model_path}"
            )
        else:
            message = (
                "File model ditemukan, tetapi loader checkpoint dan pemetaan output "
                "inferensi belum diimplementasikan. Belum ada prediksi yang dapat "
                "diterapkan ke formulir survei."
            )
        self.ai_training_status.set(message)
        messagebox.showinfo("Kesiapan Inferensi AI", message, parent=self.root)

    def cancel_survey_training(self):
        self.training_cancel_event.set()

    def open_ai_folder(self):
        try:
            os.startfile(str(SURVEY_AI_ROOT))
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def create_side_matrix(self, parent_grid, title_label, key_data, col_target):
        f_container = tk.Frame(parent_grid, bg="#f5f6fa")
        f_container.grid(row=0, column=col_target, sticky="nsew", padx=3, pady=2)
        
        tk.Label(f_container, text=f"{title_label}:", bg="#f5f6fa", font=("Arial", 8, "bold"), fg="#2c3e50").pack(anchor="w", padx=2, pady=1)
        
        options = self.options_config[key_data]
        self.buttons_ref[key_data] = {}
        
        for opt in options:
            btn = tk.Button(
                f_container, text=opt, font=("Arial", 8, "bold"), bg="#ffffff", fg="#2c3e50",
                activebackground="#00b894", activeforeground="white", relief=tk.RAISED, bd=1, pady=3,
                command=lambda k=key_data, o=opt: self.select_option_trigger(k, o)
            )
            btn.pack(fill=tk.X, padx=1, pady=2)
            self.buttons_ref[key_data][opt] = btn

    def create_perkerasan_buttons(self, parent_frame, title_label, key_data):
        f_container = tk.Frame(parent_frame, bg="#eaf2f8")
        f_container.pack(fill=tk.X, padx=10, pady=2)

        tk.Label(f_container, text=f"{title_label}:", bg="#eaf2f8", fg="#1f4e79", font=("Arial", 9, "bold")).pack(side=tk.LEFT, padx=2)
        
        f_buttons = tk.Frame(f_container, bg="#eaf2f8")
        f_buttons.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)
        
        options = self.options_config[key_data]
        self.buttons_ref[key_data] = {}

        for col, opt in enumerate(options):
            btn = tk.Button(
                f_buttons, text=opt, font=("Arial", 8, "bold"), bg="#ffffff", fg="#203040",
                activebackground="#20b486", activeforeground="white", relief=tk.RAISED, bd=1, padx=6, pady=3,
                command=lambda k=key_data, o=opt: self.select_option_trigger(k, o)
            )
            btn.grid(row=0, column=col, sticky="ew", padx=2)
            f_buttons.columnconfigure(col, weight=1)
            self.buttons_ref[key_data][opt] = btn

    def select_option_trigger(self, key_data, selected_value):
        for opt, btn in self.buttons_ref[key_data].items():
            btn.configure(bg="#ffffff", fg="#203040", relief=tk.RAISED)
        self.current_selections[key_data] = selected_value
        active_btn = self.buttons_ref[key_data][selected_value]
        active_btn.configure(bg="#00b894", fg="white", relief=tk.SUNKEN)

    def pilih_folder(self):
        folder = filedialog.askdirectory()
        if folder:
            self.ent_path.delete(0, tk.END)
            self.ent_path.insert(0, folder)
            self.photo_folder = folder

    def muat_foto(self):
        folder = self.ent_path.get().strip()
        if not folder or not os.path.exists(folder):
            messagebox.showerror("Error", "Source path folder tidak valid!")
            return

        self.photo_folder = folder
        root_key = os.path.normcase(os.path.abspath(folder))
        self.survey_data = {}
        for stored_key, data in self._stored_survey_records.items():
            try:
                stored_path = Path(stored_key).resolve()
                stored_path.relative_to(Path(folder).resolve())
            except (OSError, ValueError):
                continue
            relative_key = stored_path.relative_to(Path(folder).resolve()).as_posix()
            self.survey_data[relative_key] = data

        valid_exts = (".png", ".jpg", ".jpeg", ".bmp", ".webp")
        try:
            root_path = Path(folder)
            self.photo_list = sorted(
                path.relative_to(root_path).as_posix()
                for path in root_path.rglob("*")
                if path.is_file() and path.suffix.lower() in valid_exts
            )
        except OSError as error:
            log_application_error(f"Membaca folder {folder}", type(error), error, error.__traceback__)
            messagebox.showerror("Error", f"Gagal membaca folder:\n{error}")
            return

        self.photo_list.sort()
        if not self.photo_list:
            messagebox.showwarning("Informasi", "Tidak ditemukan file foto gambar di folder tersebut.")
            self.current_photo_idx = -1
            self.lbl_photo_info.configure(text="Foto: 0 / 0 | -")
            return

        self.current_photo_idx = 0
        self.tampilkan_foto_dan_load_data()

    def tampilkan_foto_dan_load_data(self):
        if not self.photo_list or self.current_photo_idx < 0:
            return

        filename = self.photo_list[self.current_photo_idx]
        full_path = os.path.join(self.photo_folder, filename)

        try:
            self.original_image = Image.open(full_path)
            self.zoom_fit()
        except Exception as e:
            log_application_error(f"Gagal memuat foto {full_path}", type(e), e, e.__traceback__)
            self.canvas.delete("all")
            self.canvas.create_text(200, 100, text=f"Gagal memuat gambar:\n{str(e)}", fill="white", font=("Arial", 11, "bold"))
            return

        self.lbl_photo_info.configure(text=f"Foto: {self.current_photo_idx + 1} / {len(self.photo_list)} | {filename}")

        if filename in self.survey_data:
            self.load_saved_fields_into_gui(self.survey_data[filename])
        else:
            if self.current_photo_idx > 0:
                prev_filename = self.photo_list[self.current_photo_idx - 1]
                if prev_filename in self.survey_data:
                    self.load_saved_fields_into_gui(self.survey_data[prev_filename].copy())

    def toggle_ai_mode(self):
        if self.ai_mode_enabled:
            self.ai_mode_enabled = False
            self.ai_mode_button.configure(text="MODE AI: OFF", bg="#34495e")
            self.ai_status.set("AI nonaktif")
        else:
            self.ai_mode_enabled = True
            self.ai_mode_button.configure(text="MODE AI: ON", bg="#148f77")
            self.ai_status.set("Mode AI aktif; inferensi belum terhubung")

    def render_canvas_image(self):
        if self.original_image is None:
            return

        w, h = self.original_image.size
        new_w = int(w * self.zoom_scale)
        new_h = int(h * self.zoom_scale)
        if new_w <= 0 or new_h <= 0:
            return

        resized_img = self.original_image.resize((new_w, new_h), Image.Resampling.LANCZOS)
        self.tk_photo = ImageTk.PhotoImage(resized_img)

        self.canvas.delete("all")
        self.root.update_idletasks()
        c_width = self.canvas.winfo_width()
        c_height = self.canvas.winfo_height()

        if c_width <= 1: c_width = 760
        if c_height <= 1: c_height = 570

        pos_x = max(0, (c_width - new_w) // 2)
        pos_y = max(0, (c_height - new_h) // 2)

        self.img_canvas_obj = self.canvas.create_image(pos_x, pos_y, anchor="nw", image=self.tk_photo)
        self.canvas.config(scrollregion=(0, 0, max(c_width, new_w), max(c_height, new_h)))

    def zoom_in(self):
        if self.original_image:
            self.zoom_scale *= 1.2
            self.render_canvas_image()

    def zoom_out(self):
        if self.original_image:
            self.zoom_scale /= 1.2
            self.render_canvas_image()

    def zoom_fit(self):
        if self.original_image is None:
            return

        self.root.update_idletasks()
        c_width = self.canvas.winfo_width()
        c_height = self.canvas.winfo_height()
        
        if c_width <= 1: c_width = 760
        if c_height <= 1: c_height = 570

        img_w, img_h = self.original_image.size
        ratio_w = c_width / img_w
        ratio_h = c_height / img_h
        
        self.zoom_scale = min(ratio_w, ratio_h) * 0.96
        if self.zoom_scale <= 0: 
            self.zoom_scale = 1.0

        self.render_canvas_image()

    def start_pan(self, event):
        self.canvas.scan_mark(event.x, event.y)

    def execute_pan(self, event):
        self.canvas.scan_dragto(event.x, event.y, gain=1)

    def zoom_pake_scroll_mouse(self, event):
        if self.original_image is None:
            return

        if getattr(event, "num", 0) == 4 or getattr(event, "delta", 0) > 0:
            self.zoom_scale *= 1.15
        elif getattr(event, "num", 0) == 5 or getattr(event, "delta", 0) < 0:
            self.zoom_scale /= 1.15

        if self.zoom_scale < 0.1: self.zoom_scale = 0.1
        if self.zoom_scale > 15.0: self.zoom_scale = 15.0

        self.render_canvas_image()

    def load_saved_fields_into_gui(self, data_dict):
        self.var_lebar_bahu_kiri.set(data_dict.get("lebar_bahu_kiri", "0"))
        self.var_lebar_drainase_kiri.set(data_dict.get("lebar_drainase_kiri", "0"))
        self.var_lebar_bahu_kanan.set(data_dict.get("lebar_bahu_kanan", "0"))
        self.var_lebar_drainase_kanan.set(data_dict.get("lebar_drainase_kanan", "0"))
        self.var_lebar_perkerasan.set(data_dict.get("lebar_perkerasan", "0"))

        for key in self.options_config.keys():
            val = data_dict.get(key, "")
            if val in self.options_config[key]:
                self.select_option_trigger(key, val)
            else:
                self.current_selections[key] = ""
                for opt, btn in self.buttons_ref[key].items():
                    btn.configure(bg="#ffffff", fg="#203040", relief=tk.RAISED)

    def simpan_dan_selanjutnya(self):
        if not self.photo_list or self.current_photo_idx < 0:
            return

        filename = self.photo_list[self.current_photo_idx]
        current_data = {
            "lebar_bahu_kiri": self.var_lebar_bahu_kiri.get(),
            "lebar_drainase_kiri": self.var_lebar_drainase_kiri.get(),
            "lebar_bahu_kanan": self.var_lebar_bahu_kanan.get(),
            "lebar_drainase_kanan": self.var_lebar_drainase_kanan.get(),
            "lebar_perkerasan": self.var_lebar_perkerasan.get()
        }
        current_data.update(self.current_selections)
        current_data["folder"] = self.photo_folder
        self.survey_data[filename] = current_data
        storage_key = self.photo_storage_key(filename)
        self.survey_store.save(storage_key, current_data)
        self._stored_survey_records[storage_key] = current_data.copy()

        if self.current_photo_idx < len(self.photo_list) - 1:
            self.current_photo_idx += 1
            self.tampilkan_foto_dan_load_data()
        else:
            messagebox.showinfo("Selesai", "Seluruh foto di folder ini telah dievaluasi!")

    def foto_sebelumnya(self):
        if self.current_photo_idx > 0:
            self.current_photo_idx -= 1
            self.tampilkan_foto_dan_load_data()

    def lewati_tanpa_simpan(self):
        if self.current_photo_idx < len(self.photo_list) - 1:
            self.current_photo_idx += 1
            self.tampilkan_foto_dan_load_data()

    def reset_penilaian(self):
        self.var_lebar_bahu_kiri.set("0")
        self.var_lebar_drainase_kiri.set("0")
        self.var_lebar_bahu_kanan.set("0")
        self.var_lebar_drainase_kanan.set("0")
        self.var_lebar_perkerasan.set("0")

        for key in self.options_config.keys():
            self.current_selections[key] = ""
            for opt, btn in self.buttons_ref[key].items():
                btn.configure(bg="#ffffff", fg="#203040", relief=tk.RAISED)

    def tampilkan_preview_popup(self):
        popup = tk.Toplevel(self.root)
        popup.title("Preview Data Hasil Survei")
        popup.geometry("1000x450")
        popup.configure(bg="#eaf2f8")

        tk.Label(popup, text="Daftar Data Hasil Rekaman Survei Jalan Sementara", font=("Arial", 11, "bold"), bg="#eaf2f8", pady=5).pack()

        frame_table = tk.Frame(popup)
        frame_table.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

        columns = (
            "foto", "l_bahu_k", "t_bahu_k", "k_bahu_k", "l_drain_k", "t_drain_k", "k_drain_k",
            "l_perkerasan", "t_perkerasan", "k_perkerasan",
            "l_bahu_kn", "t_bahu_kn", "k_bahu_kn", "l_drain_kn", "t_drain_kn", "k_drain_kn"
        )
        tree = ttk.Treeview(frame_table, columns=columns, show="headings")

        headers = {
            "foto": "Nama Foto", "l_bahu_k": "Lebar B.Kiri", "t_bahu_k": "Tipe B.Kiri", "k_bahu_k": "Kond B.Kiri",
            "l_drain_k": "Lebar D.Kiri", "t_drain_k": "Tipe D.Kiri", "k_drain_k": "Kond D.Kiri",
            "l_perkerasan": "Lebar Perkerasan", "t_perkerasan": "Tipe Perkerasan", "k_perkerasan": "Kond Perkerasan",
            "l_bahu_kn": "Lebar B.Kanan", "t_bahu_kn": "Tipe B.Kanan", "k_bahu_kn": "Kond B.Kanan",
            "l_drain_kn": "Lebar D.Kanan", "t_drain_kn": "Tipe D.Kanan", "k_drain_kn": "Kond D.Kanan"
        }

        for col, txt in headers.items():
            tree.heading(col, text=txt)
            tree.column(col, width=95, anchor="center")
        tree.column("foto", width=130, anchor="w")

        vsb = ttk.Scrollbar(frame_table, orient="vertical", command=tree.yview)
        hsb = ttk.Scrollbar(frame_table, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        frame_table.rowconfigure(0, weight=1)
        frame_table.columnconfigure(0, weight=1)

        for fn in self.photo_list:
            if fn in self.survey_data:
                d = self.survey_data[fn]
                tree.insert("", tk.END, values=(
                    fn, d.get("lebar_bahu_kiri",""), d.get("bahu_kiri_tipe",""), d.get("bahu_kiri_kondisi",""),
                    d.get("lebar_drainase_kiri",""), d.get("drainase_kiri_tipe",""), d.get("drainase_kiri_kondisi",""),
                    d.get("lebar_perkerasan",""), d.get("perkerasan_tipe",""), d.get("perkerasan_kondisi",""),
                    d.get("lebar_bahu_kanan",""), d.get("bahu_kanan_tipe",""), d.get("bahu_kanan_kondisi",""),
                    d.get("lebar_drainase_kanan",""), d.get("drainase_kanan_tipe",""), d.get("drainase_kanan_kondisi","")
                ))
            else:
                tree.insert("", tk.END, values=(fn,) + ("-",) * 15)

        tk.Button(popup, text="TUTUP", font=("Arial", 9, "bold"), bg="#34495e", fg="white", command=popup.destroy).pack(pady=5)

    def eksport_ke_excel(self):
        if not self.survey_data:
            messagebox.showwarning("Perhatian", "Belum ada data hasil survei yang direkam untuk dieksport!")
            return

        save_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel Files", "*.xlsx")])
        if not save_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Hasil Survei Jalan"
            ws.sheet_view.showGridLines = True

            headers_l1 = [
                "Folder Source", "Nama File Foto", 
                "BAHU JALAN KIRI", "", "", 
                "DRAINASE KIRI", "", "", 
                "PERKERASAN UTAMA", "", "", 
                "BAHU JALAN KANAN", "", "", 
                "DRAINASE KANAN", "", ""
            ]
            headers_l2 = [
                "", "", 
                "Lebar (m)", "Tipe", "Kondisi", 
                "Lebar (m)", "Tipe", "Kondisi",
                "Lebar (m)", "Tipe", "Kondisi",
                "Lebar (m)", "Tipe", "Kondisi", 
                "Lebar (m)", "Tipe", "Kondisi"
            ]

            ws.append(headers_l1)
            ws.append(headers_l2)

            ws.merge_cells('A1:A2')
            ws.merge_cells('B1:B2')
            ws.merge_cells('C1:E1')
            ws.merge_cells('F1:H1')
            ws.merge_cells('I1:K1')
            ws.merge_cells('L1:N1')
            ws.merge_cells('O1:Q1')

            f_header = XlFont(name="Arial", size=10, bold=True, color="FFFFFF")
            fill_header = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
            align_center = XlAlignment(horizontal="center", vertical="center", wrap_text=True)
            thin_border = Border(
                left=Side(style='thin', color='B0B0B0'), right=Side(style='thin', color='B0B0B0'),
                top=Side(style='thin', color='B0B0B0'), bottom=Side(style='thin', color='B0B0B0')
            )

            for row in ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=17):
                for cell in row:
                    cell.font = f_header
                    cell.fill = fill_header
                    cell.alignment = align_center
                    cell.border = thin_border

            for fn in self.photo_list:
                if fn in self.survey_data:
                    d = self.survey_data[fn]
                    ws.append([
                        self.photo_folder, fn, 
                        d.get("lebar_bahu_kiri",""), d.get("bahu_kiri_tipe",""), d.get("bahu_kiri_kondisi",""),
                        d.get("lebar_drainase_kiri",""), d.get("drainase_kiri_tipe",""), d.get("drainase_kiri_kondisi",""),
                        d.get("lebar_perkerasan",""), d.get("perkerasan_tipe",""), d.get("perkerasan_kondisi",""),
                        d.get("lebar_bahu_kanan",""), d.get("bahu_kanan_tipe",""), d.get("bahu_kanan_kondisi",""),
                        d.get("lebar_drainase_kanan",""), d.get("drainase_kanan_tipe",""), d.get("drainase_kanan_kondisi","")
                    ])
                else:
                    ws.append([self.photo_folder, fn] + ["-"] * 15)

            for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=1, max_col=17):
                for cell in row:
                    cell.font = XlFont(name="Arial", size=10)
                    cell.border = thin_border
                    if cell.column in (1, 2):
                        cell.alignment = XlAlignment(horizontal="left", vertical="center")
                    else:
                        cell.alignment = XlAlignment(horizontal="center", vertical="center")

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            wb.save(save_path)
            messagebox.showinfo("Sukses", f"Data berhasil di-eksport ke:\n{save_path}")

        except Exception as e:
            log_application_error("Ekspor ke Excel", type(e), e, e.__traceback__)
            messagebox.showerror("Ekspor Gagal", f"Terjadi kesalahan saat menyimpan file Excel:\n{e}")


def main():
    root = None
    try:
        root = tk.Tk()
        app = RoadSurveyApp(root)
        root.mainloop()
    except Exception as error:
        log_application_error("Startup Aplikasi", type(error), error, error.__traceback__)
        if root is not None:
            try:
                root.destroy()
            except tk.TclError:
                pass


if __name__ == "__main__":
    main()