import csv
import math
import os
import shutil
import tempfile
import threading
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, ttk

import piexif
from PIL import Image
from PIL.ExifTags import TAGS, GPSTAGS
import openpyxl


PHOTO_EXTENSIONS = {".jpg", ".jpeg"}


class PhotoMetadataTab(ttk.Frame):
    """Metadata EXIF/GPS tool integrated into APRJ Next."""

    def __init__(self, parent):
        super().__init__(parent, padding=10)
        self.dir_src_copy = tk.StringVar()
        self.dir_tgt_copy = tk.StringVar()
        self.overwrite_copy = tk.BooleanVar(value=False)
        self.dir_export = tk.StringVar()
        self.file_import = tk.StringVar()
        self.dir_import = tk.StringVar()
        self._busy = False
        self.subnotebook = ttk.Notebook(self)
        self.subnotebook.pack(fill="both", expand=True)
        self.copy_frame = ttk.Frame(self.subnotebook, padding=8)
        self.export_frame = ttk.Frame(self.subnotebook, padding=8)
        self.import_frame = ttk.Frame(self.subnotebook, padding=8)
        self.subnotebook.add(self.copy_frame, text="Salin Metadata")
        self.subnotebook.add(self.export_frame, text="Ekspor EXIF/GPS")
        self.subnotebook.add(self.import_frame, text="Impor GPS")
        self._build_copy_tab()
        self._build_export_tab()
        self._build_import_tab()

    def _build_copy_tab(self):
        group = ttk.LabelFrame(self.copy_frame, text="Salin metadata EXIF antar-folder", padding=10)
        group.pack(fill="x", pady=(0, 8))
        self._path_row(group, "Folder sumber metadata", self.dir_src_copy, self._browse_dir, row=0)
        self._path_row(group, "Folder foto tujuan", self.dir_tgt_copy, self._browse_dir, row=1)
        ttk.Checkbutton(
            group,
            text="Timpa foto tujuan (tanpa opsi ini, buat salinan *_with_meta.jpg)",
            variable=self.overwrite_copy,
        ).grid(row=2, column=0, columnspan=3, sticky="w", pady=4)
        ttk.Button(group, text="Salin Metadata", command=self.run_copy).grid(row=3, column=0, sticky="w", pady=5)
        self.log_copy = self._make_log("Log salin metadata", self.copy_frame)

    def _build_export_tab(self):
        group = ttk.LabelFrame(self.export_frame, text="Ekspor metadata foto", padding=10)
        group.pack(fill="x", pady=(0, 8))
        self._path_row(group, "Folder foto sumber", self.dir_export, self._browse_dir, row=0)
        ttk.Label(
            group,
            text="Membaca EXIF umum dan GPS dari file JPG/JPEG langsung di folder yang dipilih.",
        ).grid(row=1, column=0, columnspan=3, sticky="w", pady=4)
        ttk.Button(group, text="Ekstrak ke Excel/CSV", command=self.run_export).grid(row=2, column=0, sticky="w", pady=5)
        self.log_export = self._make_log("Log ekspor metadata", self.export_frame)

    def _build_import_tab(self):
        group = ttk.LabelFrame(self.import_frame, text="Impor koordinat GPS dari Excel/CSV", padding=10)
        group.pack(fill="x", pady=(0, 8))
        self._path_row(group, "File Excel/CSV", self.file_import, self._browse_data_file, row=0)
        self._path_row(group, "Folder foto tujuan", self.dir_import, self._browse_dir, row=1)
        ttk.Label(
            group,
            text="Kolom wajib: Filename, Latitude, Longitude. File yang diubah harus JPG/JPEG.",
        ).grid(row=2, column=0, columnspan=3, sticky="w", pady=4)
        ttk.Button(group, text="Tulis GPS ke Foto", command=self.run_import).grid(row=3, column=0, sticky="w", pady=5)
        self.log_import = self._make_log("Log impor GPS", self.import_frame)

    def _path_row(self, parent, label, variable, browse_command, row):
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", padx=(0, 8), pady=4)
        ttk.Entry(parent, textvariable=variable, width=72).grid(row=row, column=1, sticky="ew", pady=4)
        ttk.Button(parent, text="Pilih", command=lambda: browse_command(variable)).grid(row=row, column=2, padx=(8, 0), pady=4)
        parent.columnconfigure(1, weight=1)

    def _make_log(self, title, parent):
        frame = ttk.LabelFrame(parent, text=title, padding=5)
        frame.pack(fill="both", expand=True, pady=(0, 8))
        area = scrolledtext.ScrolledText(frame, height=7, state="disabled", font=("Consolas", 9))
        area.pack(fill="both", expand=True)
        return area

    def _browse_dir(self, variable):
        folder = filedialog.askdirectory(parent=self.winfo_toplevel())
        if folder:
            variable.set(folder)

    def _browse_data_file(self, variable):
        filename = filedialog.askopenfilename(
            parent=self.winfo_toplevel(),
            filetypes=[("Excel/CSV", "*.xlsx *.csv"), ("Excel", "*.xlsx"), ("CSV", "*.csv")],
        )
        if filename:
            variable.set(filename)

    def _log(self, widget, message):
        def append():
            try:
                widget.configure(state="normal")
                widget.insert(tk.END, message + "\n")
                widget.see(tk.END)
                widget.configure(state="disabled")
            except tk.TclError:
                pass
        self.after(0, append)

    def _start_worker(self, widget, title, worker):
        if self._busy:
            messagebox.showwarning("Proses berjalan", "Tunggu proses metadata yang sedang berjalan.", parent=self)
            return
        self._busy = True
        self._log(widget, f"--- {title} ({datetime.now():%Y-%m-%d %H:%M:%S}) ---")

        def run():
            try:
                worker()
            except Exception as error:
                self._log(widget, f"ERROR: {type(error).__name__}: {error}")
                self.after(0, lambda error=error: messagebox.showerror("Proses gagal", str(error), parent=self))
            finally:
                self.after(0, lambda: setattr(self, "_busy", False))

        threading.Thread(target=run, daemon=True).start()

    @staticmethod
    def _jpg_files(folder):
        root = Path(folder)
        if not root.is_dir():
            return []
        return sorted(
            (p for p in root.iterdir() if p.is_file() and p.suffix.lower() in PHOTO_EXTENSIONS),
            key=lambda p: p.name.casefold(),
        )

    @staticmethod
    def _decode(value):
        if isinstance(value, bytes):
            return value.decode("utf-8", errors="replace").rstrip("\x00")
        if isinstance(value, tuple):
            return str(value)
        if isinstance(value, list):
            return str(value)
        if isinstance(value, dict):
            return str(value)
        if isinstance(value, (int, float, str)) or value is None:
            return value
        try:
            return str(value)
        except Exception:
            return repr(value)

    @staticmethod
    def _coord_to_decimal(coord, ref):
        try:
            def ratio(item):
                if isinstance(item, tuple):
                    numerator, denominator = item
                    if denominator == 0:
                        raise ValueError("Penyebut koordinat GPS nol")
                    return numerator / denominator
                return float(item)
            degree, minute, second = (ratio(part) for part in coord)
            decimal = degree + minute / 60.0 + second / 3600.0
            if PhotoMetadataTab._decode(ref).upper() in ("S", "W"):
                decimal = -decimal
            return decimal
        except (TypeError, ValueError, ZeroDivisionError):
            return None

    @staticmethod
    def _decimal_to_dms(value, is_latitude):
        value = float(value)
        limit = 90 if is_latitude else 180
        if not math.isfinite(value) or abs(value) > limit:
            raise ValueError(f"Koordinat {'latitude' if is_latitude else 'longitude'} di luar rentang")
        absolute = abs(value)
        degrees = int(absolute)
        minutes_float = (absolute - degrees) * 60.0
        minutes = int(minutes_float)
        seconds_scaled = int(round((minutes_float - minutes) * 60.0 * 10000))
        if seconds_scaled >= 600000:
            seconds_scaled = 0
            minutes += 1
        if minutes >= 60:
            minutes = 0
            degrees += 1
        return ((degrees, 1), (minutes, 1), (seconds_scaled, 10000))

    def run_copy(self):
        source, target = self.dir_src_copy.get().strip(), self.dir_tgt_copy.get().strip()
        if not source or not target or not Path(source).is_dir() or not Path(target).is_dir():
            messagebox.showerror("Folder belum valid", "Pilih folder sumber dan tujuan yang tersedia.", parent=self)
            return
        if Path(source).resolve() == Path(target).resolve():
            messagebox.showerror("Folder sama", "Folder sumber dan tujuan harus berbeda.", parent=self)
            return
        overwrite = self.overwrite_copy.get()
        if overwrite and not messagebox.askyesno(
            "Konfirmasi timpa",
            "Foto tujuan dengan nama yang sama akan diubah. Lanjutkan?",
            parent=self,
        ):
            return

        def worker():
            src_files = {p.name.casefold(): p for p in self._jpg_files(source)}
            tgt_files = {p.name.casefold(): p for p in self._jpg_files(target)}
            common = sorted(src_files.keys() & tgt_files.keys())
            if not common:
                self._log(self.log_copy, "Tidak ditemukan nama JPG/JPEG yang sama di kedua folder.")
                return
            success = failed = 0
            for key in common:
                src_path, tgt_path = src_files[key], tgt_files[key]
                try:
                    exif_bytes = piexif.dump(piexif.load(str(src_path)))
                    out_path = tgt_path if overwrite else tgt_path.with_name(f"{tgt_path.stem}_with_meta{tgt_path.suffix}")
                    if out_path.exists() and not overwrite:
                        self._log(self.log_copy, f"Lewati {out_path.name}: file keluaran sudah ada.")
                        failed += 1
                        continue
                    # Use a temporary JPEG so a failed EXIF write cannot corrupt the target.
                    with tempfile.TemporaryDirectory(prefix=".aprj_meta_", dir=str(tgt_path.parent)) as temp_dir:
                        temp_path = Path(temp_dir) / f"working{tgt_path.suffix.lower()}"
                        shutil.copy2(tgt_path, temp_path)
                        piexif.insert(exif_bytes, str(temp_path))
                        if overwrite:
                            os.replace(temp_path, out_path)
                        else:
                            shutil.copy2(temp_path, out_path)
                    self._log(self.log_copy, f"Berhasil: {src_path.name} -> {out_path.name}")
                    success += 1
                except Exception as error:
                    self._log(self.log_copy, f"Gagal {src_path.name}: {error}")
                    failed += 1
            self._log(self.log_copy, f"Selesai. Berhasil: {success}; gagal/lewati: {failed}.")

        self._start_worker(self.log_copy, "Memulai salin metadata", worker)

    def run_export(self):
        source = self.dir_export.get().strip()
        if not source or not Path(source).is_dir():
            messagebox.showerror("Folder belum valid", "Pilih folder foto sumber yang tersedia.", parent=self)
            return
        save_path = filedialog.asksaveasfilename(
            parent=self.winfo_toplevel(),
            title="Simpan ekspor metadata",
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx"), ("CSV", "*.csv")],
        )
        if not save_path:
            return

        def worker():
            rows = []
            files = self._jpg_files(source)
            if not files:
                self._log(self.log_export, "Tidak ditemukan JPG/JPEG di folder tersebut.")
                return
            for path in files:
                record = {"Filename": path.name, "FilePath": str(path)}
                try:
                    data = piexif.load(str(path))
                    gps = data.get("GPS") or {}
                    lat = gps.get(piexif.GPSIFD.GPSLatitude)
                    lat_ref = gps.get(piexif.GPSIFD.GPSLatitudeRef)
                    lon = gps.get(piexif.GPSIFD.GPSLongitude)
                    lon_ref = gps.get(piexif.GPSIFD.GPSLongitudeRef)
                    record["Latitude"] = self._coord_to_decimal(lat, lat_ref) if lat and lat_ref else None
                    record["Longitude"] = self._coord_to_decimal(lon, lon_ref) if lon and lon_ref else None
                    for tag_id, value in gps.items():
                        if tag_id in (piexif.GPSIFD.GPSLatitude, piexif.GPSIFD.GPSLongitude):
                            continue
                        name = GPSTAGS.get(tag_id, str(tag_id))
                        record[f"GPS_{name}"] = self._decode(value)
                    for ifd_name in ("0th", "Exif", "1st"):
                        for tag_id, value in (data.get(ifd_name) or {}).items():
                            name = TAGS.get(tag_id, str(tag_id))
                            record[name] = self._decode(value)
                    self._log(self.log_export, f"Ekstrak: {path.name}")
                except Exception as error:
                    record["MetadataError"] = str(error)
                    self._log(self.log_export, f"Gagal {path.name}: {error}")
                rows.append(record)
            if not rows:
                return
            headers = list(dict.fromkeys(key for row in rows for key in row))
            if Path(save_path).suffix.lower() == ".csv":
                with open(save_path, "w", newline="", encoding="utf-8-sig") as stream:
                    writer = csv.DictWriter(stream, fieldnames=headers, extrasaction="ignore")
                    writer.writeheader()
                    writer.writerows(rows)
            else:
                workbook = openpyxl.Workbook()
                sheet = workbook.active
                sheet.title = "Metadata Foto"
                sheet.append(headers)
                for row in rows:
                    sheet.append([self._excel_value(row.get(key)) for key in headers])
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
                for column in sheet.columns:
                    max_len = max(len(str(cell.value or "")) for cell in column)
                    sheet.column_dimensions[column[0].column_letter].width = min(max(max_len + 2, 12), 48)
                workbook.save(save_path)
            self._log(self.log_export, f"Disimpan: {save_path}")
            self.after(0, lambda: messagebox.showinfo("Ekspor selesai", f"Metadata foto diekspor ke:\n{save_path}", parent=self))

        self._start_worker(self.log_export, "Memulai ekstraksi metadata", worker)

    @staticmethod
    def _excel_value(value):
        if isinstance(value, (dict, list, tuple)):
            return str(value)
        if isinstance(value, bytes):
            return value.decode("utf-8", errors="replace")
        return value

    @staticmethod
    def _read_table(path):
        if Path(path).suffix.lower() == ".csv":
            with open(path, "r", newline="", encoding="utf-8-sig") as stream:
                return list(csv.DictReader(stream))
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(value).strip() if value is not None else "" for value in next(iterator, ())]
            return [
                {headers[index]: row[index] for index in range(min(len(headers), len(row))) if headers[index]}
                for row in iterator
            ]
        finally:
            workbook.close()

    def run_import(self):
        data_path, target_dir = self.file_import.get().strip(), self.dir_import.get().strip()
        if not data_path or not Path(data_path).is_file() or not target_dir or not Path(target_dir).is_dir():
            messagebox.showerror("Input belum valid", "Pilih file Excel/CSV dan folder foto yang tersedia.", parent=self)
            return
        if Path(data_path).suffix.lower() not in (".csv", ".xlsx"):
            messagebox.showerror("Format tidak didukung", "Gunakan file .csv atau .xlsx.", parent=self)
            return
        try:
            rows = self._read_table(data_path)
        except Exception as error:
            messagebox.showerror("Gagal membaca data", str(error), parent=self)
            return
        if not rows or not {"Filename", "Latitude", "Longitude"}.issubset(rows[0].keys()):
            messagebox.showerror(
                "Kolom tidak lengkap",
                "File harus memiliki kolom Filename, Latitude, dan Longitude.",
                parent=self,
            )
            return
        if not messagebox.askyesno(
            "Konfirmasi penulisan GPS",
            "Koordinat GPS akan ditulis ke foto yang cocok. Setiap foto asli diproses melalui file sementara, tetapi tetap disarankan memiliki cadangan. Lanjutkan?",
            parent=self,
        ):
            return

        def worker():
            success = failed = 0
            for row in rows:
                filename = row.get("Filename")
                if filename is None or not str(filename).strip():
                    continue
                # Filename must be a basename to prevent paths escaping the selected folder.
                name = Path(str(filename).strip()).name
                if name != str(filename).strip() or Path(name).suffix.lower() not in PHOTO_EXTENSIONS:
                    self._log(self.log_import, f"Lewati nama file tidak valid: {filename}")
                    failed += 1
                    continue
                target = Path(target_dir) / name
                if not target.is_file():
                    self._log(self.log_import, f"Tidak ditemukan: {name}")
                    failed += 1
                    continue
                try:
                    lat = float(row["Latitude"])
                    lon = float(row["Longitude"])
                    if not math.isfinite(lat) or not math.isfinite(lon) or abs(lat) > 90 or abs(lon) > 180:
                        raise ValueError("Latitude/Longitude di luar rentang")
                    exif = piexif.load(str(target))
                    gps = exif.get("GPS") or {}
                    gps.update({
                        piexif.GPSIFD.GPSLatitudeRef: "S" if lat < 0 else "N",
                        piexif.GPSIFD.GPSLatitude: self._decimal_to_dms(lat, True),
                        piexif.GPSIFD.GPSLongitudeRef: "W" if lon < 0 else "E",
                        piexif.GPSIFD.GPSLongitude: self._decimal_to_dms(lon, False),
                    })
                    exif["GPS"] = gps
                    exif_bytes = piexif.dump(exif)
                    with tempfile.TemporaryDirectory(prefix=".aprj_gps_", dir=str(target.parent)) as temp_dir:
                        temp_path = Path(temp_dir) / target.name
                        shutil.copy2(target, temp_path)
                        piexif.insert(exif_bytes, str(temp_path))
                        os.replace(temp_path, target)
                    self._log(self.log_import, f"GPS ditulis: {name}")
                    success += 1
                except Exception as error:
                    self._log(self.log_import, f"Gagal {name}: {error}")
                    failed += 1
            self._log(self.log_import, f"Selesai. Berhasil: {success}; gagal/lewati: {failed}.")
            self.after(0, lambda: messagebox.showinfo(
                "Impor GPS selesai",
                f"Berhasil: {success}\nGagal/lewati: {failed}",
                parent=self,
            ))

        self._start_worker(self.log_import, "Memulai impor GPS", worker)
